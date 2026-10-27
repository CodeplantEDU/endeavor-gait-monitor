"""Editable product output example using the existing computed demonstration data."""
from pathlib import Path
import json
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'output/demo_20261027/demo-data.json').read_text(encoding='utf-8'))
R = DATA['report']
assert DATA['provenance']['type'] == 'synthetic'
assert R['probability']['value'] is None and R['ai']['analysis']['pd_probability'] is None
OUT = ROOT / 'output/doc/ENDEAVOR_70대_보행분석_출력예시_20261027.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
TEMP = ROOT / 'tmp/pdfs/product_example'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.7), Inches(.65)
sec.left_margin = sec.right_margin = Inches(.75)
sec.header_distance = sec.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Header', 'Footer']:
    style = doc.styles[name]
    style.font.name = 'Malgun Gothic'
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Malgun Gothic')
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in list(fonts.attrib):
        if 'theme' in attr.lower():
            del fonts.attrib[attr]
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.3
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].paragraph_format.space_after = Pt(10)
for name in ['Heading 1', 'Heading 2']:
    doc.styles[name].font.size = Pt(12)
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(7)
header = sec.header.paragraphs[0]
header.text = 'ENDEAVOR                                     GAIT PATTERN REPORT'
header.runs[0].font.size = Pt(9)
footer = sec.footer.paragraphs[0]
footer.text = '제품 기능 설명용 | 보고서 출력 예시                                      '
footer.runs[0].font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = '70대 1명 보행 분석 보고서 출력 예시'
doc.core_properties.subject = '제품 기능 설명용 분석 보고서'
doc.core_properties.author = 'ENDEAVOR'


def p(value, size=None, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(str(value).replace('가상', '예시'))
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    return para


def title(value, intro):
    next_page = bool(doc.paragraphs)
    para = doc.add_paragraph(value, 'Title')
    para.paragraph_format.page_break_before = next_page
    p(intro)


def heading(value):
    doc.add_paragraph(value, 'Heading 1')


def table(rows, widths, header=True):
    t = doc.add_table(rows=0, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for cell, width, value in zip(cells, widths, row):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = para.paragraph_format.space_before = Pt(4)
            para.paragraph_format.line_spacing = 1.15
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT if cell == cells[0] else WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run(str(value))
            run.font.size = Pt(9)
            run.bold = i == 0 and header
            pr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E7EDF2' if i == 0 and header else ('F7F9FA' if i % 2 else 'FFFFFF'))
            pr.append(shade)
        no_split = OxmlElement('w:cantSplit')
        t.rows[-1]._tr.get_or_add_trPr().append(no_split)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + side)
        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            e.set(qn('w:' + key), val)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    if header:
        repeat = OxmlElement('w:tblHeader')
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def num(value):
    return f'{value:.2f}'


def plot_page(page_number):
    # Reuse the vector report's graph panels so Word and PDF show identical data.
    source = TEMP / f'page-{page_number}.png'
    assert source.exists(), 'Render the product PDF before creating Word.'
    image = Image.open(source)
    w, h = image.size
    crop = image.crop((round(w * 40 / 595.276), round(h * 155 / 841.89),
                       round(w * 555 / 595.276), round(h * 650 / 841.89)))
    dest = TEMP / f'graphs-{page_number}.png'
    crop.save(dest)
    picture = doc.add_picture(str(dest), width=Inches(7))
    picture._inline.docPr.set('descr', '기준 기록과 비교 기록의 보행 분석 예시 그래프')


title('보행 패턴 분석 보고서', '70대 1명 대상 조건의 제품 보고서 출력 예시')
p('제품 기능 설명용 자료입니다. 수치와 그래프는 기능 설명을 위한 예시이며 실제 참여자의 실측 자료가 아닙니다.', bold=True)
table([
    ['문서 번호', 'EN-EXAMPLE-20261027-01'],
    ['대상 조건', '70대 1명 / 사례 설정 CASE-01'],
    ['작성일', '2026.10.27'],
    ['AI 설명 모델', R['ai']['model']],
    ['기록 구성', '동일 조건 기준 기록 3회 + 비교 기록 1회 / 양발 각 60초'],
], [1.45, 5.55], False)
heading('핵심 관찰')
p('개인 기준선 대비 보행 패턴 변화가 관찰되었습니다.', bold=True)
for key, label in [('stride_cv', '왼발 보행 주기 CV'), ('stance_cv', '왼발 접촉시간 CV'), ('front_peak_cv', '왼발 앞 FSR 피크 CV')]:
    p(f"{label}는 기준 {num(R['reference']['L'][key]['mean'])}%에서 비교 기록 {num(R['sides']['L']['metrics'][key])}%로 증가했습니다.")
p(f"양발에서 연구용 증가 기준을 넘은 항목은 {R['assessment']['flags']}개입니다. 표시 항목 수는 질병 확률이나 임상 위험도 등급이 아닙니다.")
heading('전문 의료진 상담을 고려할 조건')
p('실제 사용에서 동일 조건으로 재측정해도 변화가 지속되거나 걷기 불편과 균형 변화가 동반되면 신경과 등 전문 의료진과 상담할 수 있도록 안내합니다. 이 예시에서 특정 개인에게 진료 필요성을 판정한 것은 아닙니다.')

title('측정 조건과 정량 지표', '예시 신호에 서비스의 연구용 분석 알고리즘을 적용한 계산 결과')
table([
    ['조건', '예시 설정'],
    ['보행 방식 및 속도', '평지 직선 보행 / 1.10 m/s 입력 설정'],
    ['센서와 전송', '발당 앞뒤 FSR 2개 + MPU6500 자세각 / 20 Hz 모사'],
    ['기준 기록', '동일 장착·보정·속도 조건의 60초 기록 3회'],
    ['비교 기록 품질', '양발 각 60초 / 발별 유효 주기 52회 / 누락 0%'],
], [1.5, 5.5])
metrics = [
    ('stride_mean', '같은 발 보행 주기', 's'), ('stride_cv', '보행 주기 CV', '%'),
    ('stance_mean', 'FSR 접촉시간', 's'), ('stance_cv', '접촉시간 CV', '%'),
    ('swing_mean', 'FSR 비접촉시간', 's'), ('stance_percent', '주기 내 접촉 비율', '%'),
    ('front_peak', '앞 FSR 피크', '상대값'), ('front_peak_cv', '앞 FSR 피크 CV', '%'),
    ('rear_peak_cv', '뒤 FSR 피크 CV', '%'), ('pitch_range', '주기별 피치 범위', '도'),
]
rows = [['지표', '기준선 L / R', '비교 기록 L / R', '좌우 차이']]
for key, label, unit in metrics:
    rows.append([f'{label} ({unit})',
                 f"{num(R['reference']['L'][key]['mean'])} / {num(R['reference']['R'][key]['mean'])}",
                 f"{num(R['sides']['L']['metrics'][key])} / {num(R['sides']['R']['metrics'][key])}",
                 num(R['asymmetry'][key]) + '%'])
table(rows, [2.35, 1.6, 1.7, 1.35])
p('CV는 주기 간 표본 표준편차를 평균으로 나눈 백분율입니다. 좌우 차이는 발별 요약값 비교입니다. 두 ESP32의 시계는 하드웨어 동기화되지 않았습니다. FSR 상대값은 N·BW·kPa 단위가 아닙니다.', 9)

title('주기별 접촉 패턴', '주기를 0에서 100%로 정규화한 앞뒤 FSR 파형 비교')
p('실선은 평균, 음영은 한 사람의 주기 간 ±1 표준편차입니다. 연령 집단의 정상 범위를 나타내는 그래프가 아닙니다.', 9.5)
plot_page(3)
p('왼발 앞 FSR 피크와 파형의 일관성이 기준 기록과 다르게 설정된 예시입니다. 250점 보간은 약 50 ms인 원래 전송 간격의 정밀도를 높이지 않습니다.', 9)

title('리듬과 자세 변화', '주기별 시간과 신발 피치 및 FSR 주파수 분포')
p('발별 주기 순서, 주기에 따른 피치 변화, FSR 신호의 반복성을 함께 확인하는 화면 구성입니다.', 9.5)
plot_page(4)
p('PSD는 FSR 합산 신호의 탐색용 분석입니다. 허리 가속도를 분석한 Weiss 연구와 직접 비교하지 않습니다. 현 펌웨어가 보내지 않는 원시 가속도·각속도 기반의 RMS·Jerk 값은 제시하지 않습니다.', 9)

title('AI 보행 분석', 'GPT 6.1 Terra가 예시 데이터의 계산 결과를 바탕으로 작성한 설명')
for key, label in [('summary', '종합 요약'), ('observations', '보행 패턴 관찰'), ('quality', '기록 품질 검토'), ('limitations', '해석 시 주의사항'), ('next_steps', '재측정과 전문 의료진 상담 조건')]:
    heading(label)
    body = R['ai']['analysis'][key].replace('가상', '예시').replace('patoz2023', 'Patoz 2023')
    p(body)
p('AI 문장의 일부 용어는 출력 예시에 맞게 정리했습니다. AI 설명은 진단 결과가 아니며, 계산값과 측정 조건을 함께 검토해야 합니다.', 9)

title('판단 근거와 사용 범위', '연구 방법의 응용 범위와 예시 보고서의 해석')
heading('변화 확인 규칙')
p('20초 이상, 유효 주기 10회 이상, 동일 보정, 약 20 Hz 등 연구 품질 조건을 통과한 기록을 비교합니다. 동일 조건의 독립 기준선 3회 이상 대비 CV가 20% 이상 증가한 항목을 표시하고, 두 항목 이상이면 변화 확인을 안내합니다. 이 규칙은 팀의 연구 설정이며 임상 진단이나 낙상 위험도 기준이 아닙니다.')
heading('자료의 범위와 한계')
p('연령대는 제품 설명을 위한 대상 조건입니다. 예시 데이터가 70대 대상자의 실제 보행 분포를 대표하거나, 연령별 정확도를 검증했다는 뜻은 아닙니다. FSR 두 지점 신호는 발 전체의 압력 지도나 지면반력이 아닙니다.')
p('현재 장치에 대해 독립 검증하고 확률을 보정한 파킨슨병 예측 모델을 적용하지 않았습니다. 질병 여부나 치료 필요성은 개인의 병력과 증상, 전문 의료진의 진찰을 통해 판단해야 합니다.')
heading('참고 자료')
for ref in [
    'Patoz et al. (2023). Isolating the speed factor is crucial in gait analysis for Parkinson’s disease. Frontiers in Neuroscience. https://doi.org/10.3389/fnins.2023.1119390',
    'Weiss et al. (2011). Toward Automated, At-Home Assessment of Mobility Among Patients With Parkinson Disease, Using a Body-Worn Accelerometer. https://doi.org/10.1177/1545968311424869',
    'Anderson et al. (2026). WearGait-PD. Scientific Data 13, 440. https://doi.org/10.1038/s41597-026-06806-2',
    'Parkinson’s Foundation. Getting Diagnosed. https://www.parkinson.org/understanding-parkinsons/getting-diagnosed',
]:
    p(ref, 9)
p('재현 자료에는 고정 시드 20261027로 만든 예시 신호, 분석 계산 결과, 기존 OpenAI API 응답이 포함됩니다. 실제 참여자의 개인정보나 의료진 서명은 포함하지 않았습니다.', 9)

assert len(doc.inline_shapes) == 2
for element in doc.styles.element.xpath('.//w:pBdr'):
    element.getparent().remove(element)
doc.save(OUT)
print(OUT)
