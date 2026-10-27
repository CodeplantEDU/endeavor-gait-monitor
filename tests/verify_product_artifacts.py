"""Check delivered examples and unchanged firmware package without hardware I/O."""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
from pypdf import PdfReader
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
data = json.loads((ROOT / 'output/demo_20261027/demo-data.json').read_text(encoding='utf-8'))
assert data['provenance']['type'] == 'synthetic'
pdf = ROOT / 'output/pdf/ENDEAVOR_70대_보행분석_출력예시_20261027.pdf'
word = ROOT / 'output/doc/ENDEAVOR_70대_보행분석_출력예시_20261027.docx'
reader = PdfReader(pdf)
assert len(reader.pages) == 6
text = '\n'.join(p.extract_text() for p in reader.pages)
assert '70대1명' in ''.join(text.split()) and '실측자료가아닙니다' in ''.join(text.split())
assert '가상' not in text
assert all('보고서 출력 예시' in p.extract_text() for p in reader.pages)
doc = Document(word)
word_text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
assert '실측 자료가 아닙니다' in word_text
assert len(doc.inline_shapes) == 2
rendered = PdfReader(next((ROOT / 'tmp/pdfs/product_example_docx').glob('*.pdf')))
assert len(rendered.pages) == 6
assert all(len(page.extract_text()) > 200 for page in rendered.pages)
assert not doc.styles['Title'].element.xpath('.//w:pBdr')
for key in ['stride_cv', 'stance_cv', 'front_peak_cv']:
    value = f"{data['report']['sides']['L']['metrics'][key]:.2f}"
    assert value in text and value in word_text
with ZipFile(ROOT / 'output/firmware_basic/ESP32_MPU6500_FSR_BLE_기본펌웨어.zip') as archive:
    files = [n for n in archive.namelist() if not n.endswith('/')]
    assert len(files) == 2
    for file in files:
        name = file.replace('\\', '/').split('/')[-1]
        assert archive.read(file) == (ROOT / 'firmware/fsr_mpu_ble' / name).read_bytes()
expected = {
    'fsr_mpu_ble.ino': 'e59d540e88b754f685d4751dc77d117cb5823e8cdd4dbc977e9bc7d50c3a66e6',
    'OrientationFilter.h': 'c323062782cc4fa58c3e311d51b66add31367954efcb64217d755ba59e467d49',
}
for name, digest in expected.items():
    assert hashlib.sha256((ROOT / 'firmware/fsr_mpu_ble' / name).read_bytes()).hexdigest() == digest
print('PASS: 6-page example, measured-data distinction, JSON metrics, Word figures, unchanged firmware ZIP')
